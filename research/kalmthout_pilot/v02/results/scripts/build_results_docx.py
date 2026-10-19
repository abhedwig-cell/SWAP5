from pathlib import Path
import re,textwrap
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
R=Path(__file__).resolve().parents[1];source=R/'Kalmthout_V02_results_and_SYMPHONICA.md';dest=R/'Kalmthout_V02_results_and_SYMPHONICA.docx'
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption'):
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(6);doc.styles['Normal'].paragraph_format.line_spacing=1.08
for n,size in [('Title',25),('Heading 1',16),('Heading 2',12)]:
 s=doc.styles[n];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.space_before=Pt(13);s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=True
doc.styles['Subtitle'].font.size=Pt(13);doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.italic=False;doc.styles['Caption'].paragraph_format.space_after=Pt(9);doc.styles['Caption'].paragraph_format.keep_with_next=False
header=sec.header.paragraphs[0];header.text='SYMPHONICA   |   Kalmthout SWAP V02   |   Results and conversion';header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor.from_string('555555')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('Provisional regional results   |   19 October 2026   |   ').font.size=Pt(8);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='Kalmthoutse Heide SWAP results and SYMPHONICA products';doc.core_properties.author='SYMPHONICA';doc.core_properties.subject='Thematic results, temporal and spatial patterns and explicit conversion of provisional SWAP431 output'
def table(rows):
 n=len(rows[0]);widths={2:[5.5,11.7],3:[4.1,6.3,6.8],4:[7.4,2.3,3.75,3.75],5:[3.7,3.4,3.5,3.3,3.3],6:[3.4,2.6,2.7,2.7,2.7,3.1],7:[2.7,2.4,2.4,2.4,2.4,2.4,2.5],8:[2.8,1.8,1.8,1.8,2,2,2.4,2.6]}[n]
 if rows[0][0]=='Actual centre cm below surface':widths=[4.3,4.3,4.3,4.3]
 if rows[0][0]=='Analysis population':widths=[9.2,1.5,3.25,3.25]
 if rows[0][0]=='Candidate file':widths=[9.1,2.0,1.5,1.5,3.1]
 if rows[0][0]=='Year' and n==7:widths=[1.1,4.0,2.1,2.1,2.6,2.6,2.7]
 if rows[0][0]=='Year' and rows[0][-1]=='Mean GPP µmol CO2 m⁻² s⁻¹':widths=[1.4,2.5,3.3,4,6]
 if rows[0][0]=='Year' and n==4:widths=[1.8,3.2,6.2,6]
 t=doc.add_table(rows=0,cols=n);t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Cm(w)
 for k,row in enumerate(rows):
  cells=t.add_row().cells
  for j,value in enumerate(row):
   cells[j].width=Cm(widths[j]);cells[j].vertical_alignment=1;p=cells[j].paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.0
   r=p.add_run(value);r.font.size=Pt(8.1 if n>=7 else 8.8);r.bold=k==0
   if len(rows)<=7 and k<len(rows)-1:p.paragraph_format.keep_with_next=True
   pr=cells[j]._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7ECEF' if k==0 else ('F5F6F7' if k%2==0 else 'FFFFFF'));pr.append(sh)
  trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if k==0:trpr.append(OxmlElement('w:tblHeader'))
 b=OxmlElement('w:tblBorders')
 for side in ('top','bottom','left','right','insideH','insideV'):
  el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:color'),'D9D9D9');el.set(qn('w:sz'),'4');b.append(el)
 t._tbl.tblPr.append(b);doc.add_paragraph().paragraph_format.space_after=Pt(0)
lines=source.read_text().splitlines();i=0
while i<len(lines):
 line=lines[i]
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   row=[s.strip() for s in lines[i].strip('|').split('|')]
   if not all(re.fullmatch(r':?-+:?',s) for s in row):rows.append(row)
   i+=1
  table(rows);continue
 if line.startswith('!['):
  target=re.search(r'\]\((.*?)\)',line)[1];p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
  width=17.2;im=Image.open(R/target);height=width*im.height/im.width
  if height>18:width*=18/height
  p.add_run().add_picture(str(R/target),width=Cm(width))
  drawing=p._p.xpath('.//wp:docPr')
  if drawing:drawing[0].set('descr',line[2:line.index(']')])
 elif line.startswith('Figure '):doc.add_paragraph(line,'Caption')
 elif line.startswith('# '):doc.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
 elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
 elif line.startswith('```'):
  i+=1
  while i<len(lines) and not lines[i].startswith('```'):
   for segment in textwrap.wrap(lines[i],width=95,break_long_words=True,break_on_hyphens=False) or ['']:
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);r=p.add_run(segment);r.font.name='DejaVu Sans Mono';r.font.size=Pt(7.2)
   i+=1
 elif line:
  doc.add_paragraph(line,'Subtitle' if i==2 else 'Normal')
 i+=1
doc.save(dest);print(dest)
